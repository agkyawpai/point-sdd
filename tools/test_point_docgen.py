"""Run from the repo root: .venv\\Scripts\\python -m unittest tools/test_point_docgen.py"""

import json
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

TOOL = Path(__file__).with_name("point-docgen.py")

PROPOSAL = """## Why

Customers wait at the counter because barbers cannot see who has booked.

## What Changes

- Add an appointment queue view for barbers
- Add a no-show flag to appointments
- Add an appointment queue view for barbers

## Impact

Front desk only.
"""

DESIGN = """## Context

Small shop.

## Decisions

### Queue is ordered by start time

Walk-ins are slotted after booked customers.

**Decision 2 - Mark no-shows after 15 minutes.**

A late customer is flagged so the next one can be served.

## Risks / Trade-offs

None.
"""

SPEC = """## ADDED Requirements

### Requirement: Barber sees the queue

#### Scenario: Booked customer appears in the queue
- **WHEN** a customer books a 10:00 slot
- **THEN** the barber queue SHALL show the customer

#### Scenario: Customer is 15 minutes late
- **WHEN** 15 minutes pass after the slot start
- **THEN** the appointment SHALL be flagged as no-show
"""

TASKS = """## 1. Build

- [ ] 1.1 Add queue screen

## 2. Verification

- [ ] 2.1 Book a slot and open the queue
- [x] 2.2 Wait past the grace period and check the flag
"""


def full_doc(image_path=""):
    return {
        "meta": {"title": "Appointment Queue", "outputSuffix": "Appointment Queue"},
        "overview": {"sections": [
            {"title": "1. Why the change", "body": "Customers wait."},
            {"title": "2. What is changed", "body": "• Queue view\n• No-show flag"},
            {"title": "3. How the process works", "body": "1. Book\n2. Queue\n3. Serve"},
            {"title": "4. How to test", "body": "Conditions:\n• Booked\n\nSteps:\n1. Open queue"},
        ]},
        "uiux": {"note": "Screens below.", "images": [{"caption": "queue", "path": image_path}]},
        "diagrams": {"title": "Queue Flow", "instruction": "Render it.",
                     "mermaid": "flowchart TD\n A[Book] --> B[Queue]", "renderedImage": ""},
        "detailedDesign": {
            "decisions": [{"process": "Ordering", "contents": "1. By start time"},
                          {"process": "No-show", "contents": "1. After 15 minutes"}],
            "apiStructure": {
                "present": True,
                "route": "GET /api/appointments/queue",
                "responseSummary": "List of queued appointments.",
                "requestParams": [{"field": "date", "type": "string", "required": "yes",
                                   "desc": "Day", "isNew": False},
                                  {"field": "barber_id", "type": "int", "required": "no",
                                   "desc": "Filter", "isNew": True}],
                "responseRows": [{"field": "data[]", "isSection": True},
                                 {"field": "no_show", "type": "bool", "when": "always",
                                  "desc": "Flag", "isNew": True}],
            },
        },
        "databaseDesign": {"present": True, "tables": [{
            "name": "appointments", "description": "Add no-show flag",
            "columns": [{"name": "id", "type": "bigint", "null": "NO", "isNew": False},
                        {"name": "no_show", "type": "tinyint(1)", "null": "NO", "default": "0",
                         "desc": "Late customer", "isNew": True}],
        }]},
    }


class DocGenTests(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, self.tmp, ignore_errors=True)
        self.change = self.tmp / "openspec" / "changes" / "add-appointment-queue"
        (self.change / "specs" / "appointments").mkdir(parents=True)
        (self.change / "proposal.md").write_text(PROPOSAL, encoding="utf-8")
        (self.change / "design.md").write_text(DESIGN, encoding="utf-8")
        (self.change / "tasks.md").write_text(TASKS, encoding="utf-8")
        (self.change / "specs" / "appointments" / "spec.md").write_text(SPEC, encoding="utf-8")
        self.out_dir = self.tmp / "work" / "Design Docs"
        self.out_dir.mkdir(parents=True)

    def run_tool(self, *args, check=True):
        return subprocess.run([sys.executable, str(TOOL), *map(str, args)],
                              capture_output=True, text=True, check=check)

    def build(self, doc):
        src = self.tmp / "work" / "temp_doc.json"
        src.write_text(json.dumps(doc, ensure_ascii=False), encoding="utf-8")
        out = self.out_dir / f"{doc['meta']['outputSuffix']} - Documentation.xlsx"
        self.run_tool("--input", src, "--output", out)
        return load_workbook(out)

    # -- workbook from reviewed JSON

    def test_full_workbook_has_five_sheets(self):
        wb = self.build(full_doc())
        self.assertEqual(wb.sheetnames, ["Overview", "UIUX", "Diagrams", "Detailed Design", "Database Design"])

    def test_overview_title_and_sections(self):
        ws = self.build(full_doc())["Overview"]
        self.assertEqual(ws["B1"].value, "Appointment Queue - Change Documentation")
        self.assertIn("B1:H1", [str(r) for r in ws.merged_cells.ranges])
        self.assertEqual(ws["B3"].value, "1. Why the change")
        self.assertEqual(ws["B4"].value, "Customers wait.")
        self.assertEqual(ws["B12"].value, "4. How to test")
        self.assertEqual(ws["B1"].fill.fgColor.rgb, "FF1F4E78")

    def test_uiux_and_diagrams(self):
        wb = self.build(full_doc(image_path=str(self.tmp / "missing.png")))
        ui = wb["UIUX"]
        self.assertEqual(ui["B1"].value, "Appointment Queue - UI/UX")
        self.assertEqual(ui["B3"].value, "Screens below.")
        self.assertEqual(ui["B5"].value, "queue")
        dg = wb["Diagrams"]
        self.assertEqual(dg["B1"].value, "Queue Flow")
        self.assertEqual(dg["B5"].value, "Mermaid source:")
        self.assertTrue(dg["B6"].value.startswith("flowchart TD"))
        self.assertIn("B6:H11", [str(r) for r in dg.merged_cells.ranges])
        self.assertEqual(dg["B14"].value, "Insert rendered diagram PNG below:")

    def test_detailed_design_layout(self):
        ws = self.build(full_doc())["Detailed Design"]
        self.assertEqual([ws["A3"].value, ws["B3"].value, ws["C3"].value], ["No", "Process", "Processing Contents"])
        self.assertEqual([ws["A4"].value, ws["B4"].value, ws["C4"].value], [1, "Ordering", "1. By start time"])
        self.assertEqual(ws["A6"].value, "API Structure")
        self.assertEqual([ws["A7"].value, ws["B7"].value], ["Route", "GET /api/appointments/queue"])
        self.assertEqual(ws["A9"].value, "Request Parameters")
        self.assertEqual([ws.cell(10, c).value for c in (1, 2, 3, 4)], ["Field", "Type", "Required", "Description"])
        self.assertEqual(ws["A12"].value, "barber_id")
        self.assertEqual(ws["A12"].fill.fgColor.rgb, "FFFFF2CC")
        self.assertTrue(ws["A12"].font.bold)
        self.assertEqual(ws["A14"].value, "Response Structure")
        self.assertEqual(ws["A15"].value, "List of queued appointments.")
        self.assertEqual(ws["C16"].value, "When present")
        self.assertEqual(ws["A17"].value, "data[]")
        self.assertIn("A17:I17", [str(r) for r in ws.merged_cells.ranges])
        self.assertEqual(ws["A18"].value, "no_show")

    def test_database_design_layout(self):
        ws = self.build(full_doc())["Database Design"]
        self.assertEqual([ws["A1"].value, ws["B1"].value, ws["C1"].value], ["1", "Table Name", "Change Description"])
        self.assertEqual([ws["B2"].value, ws["C2"].value], ["appointments", "Add no-show flag"])
        self.assertEqual([c.value for c in ws[3]][:8],
                         ["#", "Name", "Type", "Null", "Default", "Extra", "Comments", "Description"])
        self.assertEqual(ws["B5"].value, "no_show")
        self.assertEqual(ws["A5"].fill.fgColor.rgb, "FFFFF2CC")
        self.assertNotEqual(ws["A4"].fill.fgColor.rgb, "FFFFF2CC")

    def test_optional_parts_are_skipped(self):
        doc = full_doc()
        doc["databaseDesign"] = {"present": False}
        doc["detailedDesign"]["apiStructure"] = {"present": False}
        wb = self.build(doc)
        self.assertNotIn("Database Design", wb.sheetnames)
        self.assertIsNone(wb["Detailed Design"]["A6"].value)

    def test_errors(self):
        res = self.run_tool("--input", self.tmp / "nope.json", "--output", self.out_dir / "x.xlsx", check=False)
        self.assertEqual(res.returncode, 1)
        self.assertIn("Input file not found", res.stdout)
        src = self.tmp / "doc.json"
        src.write_text("{}", encoding="utf-8")
        res = self.run_tool("--input", src, "--output", self.tmp / "no-dir" / "x.xlsx", check=False)
        self.assertEqual(res.returncode, 1)
        self.assertIn("Output directory does not exist", res.stdout)

    # -- draft from a change folder, then build

    def test_draft_from_change_folder_then_build(self):
        draft_path = self.tmp / "work" / "temp_doc.json"
        self.run_tool("--draft", self.change, "--json-out", draft_path)
        draft = json.loads(draft_path.read_text(encoding="utf-8"))
        sections = {s["title"]: s["body"] for s in draft["overview"]["sections"]}
        self.assertIn("barbers cannot see", sections["1. Why the change"])
        self.assertEqual(sections["2. What is changed"].count("appointment queue view"), 1)
        self.assertIn("1. Queue is ordered by start time", sections["3. How the process works"])
        self.assertIn("• Customer is 15 minutes late", sections["4. How to test"])
        self.assertIn("2. Wait past the grace period", sections["4. How to test"])
        processes = [d["process"] for d in draft["detailedDesign"]["decisions"]]
        self.assertEqual(processes, ["Queue is ordered by start time", "Decision 2 - Mark no-shows after 15 minutes"])
        self.assertEqual(draft["meta"]["outputSuffix"], "add-appointment-queue")
        self.assertEqual(draft["meta"]["missingArtifacts"]["add-appointment-queue"], [])

        wb = self.build(draft)
        self.assertEqual(wb.sheetnames, ["Overview", "UIUX", "Diagrams", "Detailed Design"])
        self.assertEqual(wb["Detailed Design"]["B5"].value, "Decision 2 - Mark no-shows after 15 minutes")
        self.assertTrue((self.out_dir / "add-appointment-queue - Documentation.xlsx").exists())

    def test_draft_reports_missing_artifacts(self):
        (self.change / "tasks.md").unlink()
        (self.change / "design.md").unlink()
        res = self.run_tool("--draft", self.change, "--json-out", self.tmp / "d.json")
        self.assertIn("missing design.md, tasks.md", res.stdout)


if __name__ == "__main__":
    unittest.main()
