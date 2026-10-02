"""
point-docgen.py - build the change documentation workbook for the Point Barbershop System.

Build the workbook from reviewed JSON (the normal path, driven by the
point-generate-docs skill):

    .venv\\Scripts\\python tools\\point-docgen.py --input work\\temp_doc.json \\
        --output "work\\Design Docs\\<suffix> - Documentation.xlsx"

Optionally draft that JSON straight from one or more OpenSpec change folders
(a starting point for the skill; the user still reviews it):

    .venv\\Scripts\\python tools\\point-docgen.py --draft openspec\\changes\\<name> \\
        --json-out work\\temp_doc.json

Sheets: Overview, UIUX, Diagrams, Detailed Design and, when the change touches
the database, Database Design. Only dependency: openpyxl.
"""
import argparse
import json
import os
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# --------------------------------------------------------------------------- style

COLOR = {
    "header": "FF1F4E78",   # title bars and table headers
    "section": "FFD9EAF7",  # section dividers
    "soft": "FFEAF3F8",     # notes, route row, table description row
    "code": "FFF8FAFC",     # mermaid source block
    "new": "FFFFF2CC",      # added fields / columns
    "white": "FFFFFFFF",
    "ink": "FF1F2937",
    "muted": "FF6B7280",
}


def fill(key):
    return PatternFill("solid", fgColor=COLOR[key])


NO_FILL = PatternFill(fill_type=None)

FONT = {
    "title": Font(name="Aptos Display", size=16, bold=True, color=COLOR["white"]),
    "section": Font(name="Aptos", size=12, bold=True, color=COLOR["ink"]),
    "th": Font(name="Aptos", size=10, bold=True, color=COLOR["white"]),
    "body": Font(name="Aptos", size=10, color=COLOR["ink"]),
    "strong": Font(name="Aptos", size=10, bold=True, color=COLOR["ink"]),
    "hint": Font(name="Aptos", size=11, color=COLOR["muted"]),
    "code": Font(name="Consolas", size=10, color=COLOR["ink"]),
    "table": Font(name="Aptos", size=12, bold=True, color=COLOR["white"]),
}

ALIGN = {
    "top": Alignment(horizontal="left", vertical="top", wrap_text=True),
    "mid": Alignment(horizontal="left", vertical="center", wrap_text=True),
    "center": Alignment(horizontal="center", vertical="center", wrap_text=True),
    "center_top": Alignment(horizontal="center", vertical="top", wrap_text=True),
}

THIN = Side(style="thin")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".gif"}
DRAWIO_NOTE = ("Mermaid source is provided below. Import it into draw.io (https://app.diagrams.net) "
               "or render it at https://mermaid.live, export a PNG and insert it in the rows below.")


def put(ws, row, col, value, bg, font, align):
    cell = ws.cell(row=row, column=col)
    cell.value = value
    cell.fill = bg
    cell.font = font
    cell.alignment = align
    cell.border = BOX
    return cell


def put_span(ws, row, first, last, value, bg, font, align, height=None):
    """Merge first..last on one row, write the value, box every cell in the span."""
    ws.merge_cells(start_row=row, start_column=first, end_row=row, end_column=last)
    put(ws, row, first, value, bg, font, align)
    for col in range(first + 1, last + 1):
        ws.cell(row=row, column=col).border = BOX
    if height:
        ws.row_dimensions[row].height = height


def text_height(text, minimum, per_line=15, extra=1):
    return max(minimum, per_line * (str(text or "").count("\n") + 1 + extra))


def narrow_page(ws, title):
    """Shared frame for the single-column sheets: gutter / wide B / gutter + title bar."""
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 115
    ws.column_dimensions["C"].width = 3
    put_span(ws, 1, 2, 8, title, fill("header"), FONT["title"], ALIGN["mid"], height=32)


def image_class():
    try:
        from openpyxl.drawing.image import Image  # needs Pillow at runtime
        return Image
    except Exception:
        return None


# --------------------------------------------------------------------------- sheets

def build_overview(ws, overview, title):
    narrow_page(ws, title)
    row = 3
    for section in overview.get("sections", []):
        put(ws, row, 2, section.get("title", ""), fill("section"), FONT["section"], ALIGN["mid"])
        ws.row_dimensions[row].height = 24
        body = section.get("body", "")
        put(ws, row + 1, 2, body, fill("white"), FONT["body"], ALIGN["top"])
        ws.row_dimensions[row + 1].height = text_height(body, 72)
        row += 3


def build_uiux(ws, uiux, title):
    narrow_page(ws, title)
    put_span(ws, 3, 2, 8, uiux.get("note", ""), fill("soft"), FONT["hint"], ALIGN["top"], height=28)

    images = uiux.get("images", [])
    image_cls = image_class() if images else None
    target_width = int(115 * 7.5)  # approx pixel width of column B
    row = 5
    for entry in images:
        put(ws, row, 2, entry.get("caption", ""), fill("section"), FONT["section"], ALIGN["mid"])
        ws.row_dimensions[row].height = 24
        row += 1
        path = entry.get("path", "")
        if not (image_cls and path and os.path.exists(path)):
            row += 1
            continue
        try:
            pic = image_cls(path)
            if pic.width:
                ratio = target_width / pic.width
                pic.width, pic.height = int(pic.width * ratio), int(pic.height * ratio)
            ws.add_image(pic, f"B{row}")
            rows_used = max(1, int((pic.height or 200) * 0.75 / 15) + 1)
            for r in range(row, row + rows_used):
                ws.row_dimensions[r].height = 15
            row += rows_used + 1
        except Exception as exc:
            print(f"[WARN] Skipped image {path}: {exc}")
            row += 1


def build_diagrams(ws, diagrams, title):
    narrow_page(ws, diagrams.get("title") or title)
    put_span(ws, 3, 2, 8, diagrams.get("instruction", ""), fill("soft"), FONT["body"], ALIGN["top"], height=42)
    put_span(ws, 5, 2, 8, "Mermaid source:", fill("section"), FONT["section"], ALIGN["mid"], height=24)

    ws.merge_cells("B6:H11")
    put(ws, 6, 2, diagrams.get("mermaid", ""), fill("code"), FONT["code"], ALIGN["top"])
    for r in range(6, 12):
        ws.row_dimensions[r].height = 18
        for c in range(2, 9):
            ws.cell(row=r, column=c).border = BOX

    put_span(ws, 14, 2, 8, "Insert rendered diagram PNG below:", fill("section"), FONT["section"],
             ALIGN["mid"], height=24)
    rendered = diagrams.get("renderedImage", "")
    if rendered and os.path.exists(rendered):
        image_cls = image_class()
        try:
            ws.add_image(image_cls(rendered), "B15")
        except Exception as exc:
            print(f"[WARN] Could not insert diagram image: {exc}")


def _table_header(ws, row, labels):
    """Three single header cells in A..C, then the last label spanning D..I."""
    ws.row_dimensions[row].height = 24
    for col, label in enumerate(labels[:3], start=1):
        put(ws, row, col, label, fill("header"), FONT["th"], ALIGN["center"])
    put_span(ws, row, 4, 9, labels[3], fill("header"), FONT["th"], ALIGN["center"])


def _field_row(ws, row, values, is_new):
    bg = fill("new") if is_new else NO_FILL
    font = FONT["strong"] if is_new else FONT["body"]
    ws.row_dimensions[row].height = 32
    for col, value in enumerate(values[:3], start=1):
        put(ws, row, col, value, bg, font, ALIGN["top"])
    put_span(ws, row, 4, 9, values[3], bg, font, ALIGN["top"])


def build_detailed_design(ws, design):
    ws.sheet_view.showGridLines = False
    for col, width in (("A", 22), ("C", 16), ("D", 60), ("E", 3)):
        ws.column_dimensions[col].width = width
    put_span(ws, 1, 1, 9, "Detailed Design", fill("header"), FONT["title"], ALIGN["mid"], height=32)

    ws.row_dimensions[3].height = 24
    put(ws, 3, 1, "No", fill("header"), FONT["th"], ALIGN["center"])
    put(ws, 3, 2, "Process", fill("header"), FONT["th"], ALIGN["center"])
    put_span(ws, 3, 3, 9, "Processing Contents", fill("header"), FONT["th"], ALIGN["center"])

    decisions = design.get("decisions", [])
    for no, decision in enumerate(decisions, start=1):
        row = 3 + no
        contents = decision.get("contents", "")
        ws.row_dimensions[row].height = text_height(contents, 74, extra=2)
        put(ws, row, 1, no, NO_FILL, FONT["body"], ALIGN["center"])
        put(ws, row, 2, decision.get("process", ""), NO_FILL, FONT["body"], ALIGN["center"])
        put_span(ws, row, 3, 9, contents, NO_FILL, FONT["body"], ALIGN["top"])

    api = design.get("apiStructure", {})
    if not api.get("present"):
        return

    row = 4 + len(decisions)
    put_span(ws, row, 1, 9, "API Structure", fill("section"), FONT["section"], ALIGN["mid"], height=24)
    row += 1
    ws.row_dimensions[row].height = 35
    put(ws, row, 1, "Route", fill("soft"), FONT["strong"], ALIGN["center_top"])
    put_span(ws, row, 2, 9, api.get("route", ""), fill("soft"), FONT["body"], ALIGN["top"])

    row += 2
    put_span(ws, row, 1, 9, "Request Parameters", fill("section"), FONT["section"], ALIGN["mid"], height=24)
    row += 1
    _table_header(ws, row, ["Field", "Type", "Required", "Description"])
    params = api.get("requestParams", [])
    for i, p in enumerate(params, start=1):
        _field_row(ws, row + i,
                   [p.get("field", ""), p.get("type", ""), p.get("required", ""), p.get("desc", "")],
                   p.get("isNew", False))

    row += len(params) + 2
    put_span(ws, row, 1, 9, "Response Structure", fill("section"), FONT["section"], ALIGN["mid"], height=24)
    row += 1
    put_span(ws, row, 1, 9, api.get("responseSummary", ""), fill("soft"), FONT["body"], ALIGN["top"], height=52)
    row += 1
    _table_header(ws, row, ["Field", "Type", "When present", "Description"])
    for i, r in enumerate(api.get("responseRows", []), start=1):
        if r.get("isSection"):
            put_span(ws, row + i, 1, 9, r.get("field", ""), fill("section"), FONT["strong"],
                     ALIGN["center"], height=32)
            continue
        _field_row(ws, row + i,
                   [r.get("field", ""), r.get("type", ""), r.get("when", ""), r.get("desc", "")],
                   r.get("isNew", False))


DB_COLUMNS = ["#", "Name", "Type", "Null", "Default", "Extra", "Comments", "Description"]
DB_KEYS = ["name", "type", "null", "default", "extra", "comments", "desc"]


def build_database_design(ws, database):
    ws.sheet_view.showGridLines = False
    for col, width in zip("ABCDEFGH", (6, 22, 20, 9, 16, 23, 25, 50)):
        ws.column_dimensions[col].width = width

    row = 1
    for t_no, table in enumerate(database.get("tables", []), start=1):
        ws.row_dimensions[row].height = 28
        put(ws, row, 1, str(t_no), fill("header"), FONT["table"], ALIGN["mid"])
        put(ws, row, 2, "Table Name", fill("header"), FONT["table"], ALIGN["mid"])
        put_span(ws, row, 3, 8, "Change Description", fill("header"), FONT["table"], ALIGN["mid"])
        row += 1

        ws.row_dimensions[row].height = 46
        put(ws, row, 1, "", fill("soft"), FONT["body"], ALIGN["top"])
        put(ws, row, 2, table.get("name", ""), fill("soft"), FONT["body"], ALIGN["top"])
        put_span(ws, row, 3, 8, table.get("description", ""), fill("soft"), FONT["body"], ALIGN["top"])
        row += 1

        ws.row_dimensions[row].height = 24
        for col, label in enumerate(DB_COLUMNS, start=1):
            put(ws, row, col, label, fill("header"), FONT["th"], ALIGN["center"])
        row += 1

        for c_no, column in enumerate(table.get("columns", []), start=1):
            is_new = column.get("isNew", False)
            bg = fill("new") if is_new else NO_FILL
            font = FONT["strong"] if is_new else FONT["body"]
            ws.row_dimensions[row].height = 34
            values = [str(c_no)] + [column.get(k, "") for k in DB_KEYS]
            for col, value in enumerate(values, start=1):
                put(ws, row, col, value, bg, font, ALIGN["top"])
            row += 1
        row += 1  # blank spacer row between tables


def build_workbook(data):
    title = data.get("meta", {}).get("title") or "Change Documentation"
    wb = Workbook()
    ws = wb.active
    ws.title = "Overview"
    build_overview(ws, data.get("overview", {}), f"{title} - Change Documentation")
    build_uiux(wb.create_sheet("UIUX"), data.get("uiux", {}), f"{title} - UI/UX")
    build_diagrams(wb.create_sheet("Diagrams"), data.get("diagrams", {}), f"{title} - Diagrams")
    build_detailed_design(wb.create_sheet("Detailed Design"), data.get("detailedDesign", {}))
    database = data.get("databaseDesign", {})
    if database.get("present"):
        build_database_design(wb.create_sheet("Database Design"), database)
    return wb


# --------------------------------------------------------------------------- draft from a change folder

def _read(path):
    return path.read_text(encoding="utf-8") if path.is_file() else ""


def _section(markdown, heading):
    """Body of the first '## <heading>' section (case-insensitive), up to the next '## '."""
    match = re.search(rf"^##\s+{re.escape(heading)}\s*$(.*?)(?=^##\s|\Z)", markdown,
                      flags=re.MULTILINE | re.DOTALL | re.IGNORECASE)
    return match.group(1).strip() if match else ""


def _bullets(text):
    items = []
    for line in text.splitlines():
        m = re.match(r"^\s*[-*]\s+(.*)", line)
        if m:
            items.append(m.group(1).strip())
    return items


def _decisions(design_md):
    """Split the Decisions section into (title, body) pairs.

    Accepts '### Title' sub-headings or a bold lead line such as '**Decision 1 - Title.**'."""
    body = _section(design_md, "Decisions")
    if not body:
        return []
    pattern = re.compile(r"^(?:###\s+(.+)|\*\*(.+?)\*\*\s*)$", flags=re.MULTILINE)
    marks = list(pattern.finditer(body))
    if not marks:
        return [("Decisions", body)]
    found = []
    for i, m in enumerate(marks):
        end = marks[i + 1].start() if i + 1 < len(marks) else len(body)
        title = (m.group(1) or m.group(2)).strip().rstrip(".")
        found.append((title, body[m.end():end].strip()))
    return found


def _scenarios(spec_md):
    return [m.group(1).strip() for m in re.finditer(r"^####\s+Scenario:\s*(.+)$", spec_md, re.MULTILINE)]


def _verification_steps(tasks_md):
    steps = []
    for m in re.finditer(r"^##[^\n]*verif[^\n]*$(.*?)(?=^##\s|\Z)", tasks_md,
                         flags=re.MULTILINE | re.DOTALL | re.IGNORECASE):
        for line in m.group(1).splitlines():
            t = re.match(r"^\s*-\s*\[[ xX]\]\s*(?:[\d.]+\s+)?(.*)", line)
            if t:
                steps.append(t.group(1).strip())
    return steps


def _dedupe(items):
    seen, out = set(), []
    for item in items:
        if item not in seen:
            seen.add(item)
            out.append(item)
    return out


def draft_from_changes(change_dirs):
    """Draft the review JSON from OpenSpec change folders. Content is a starting point only."""
    whys, whats, decisions, scenarios, steps, missing = [], [], [], [], [], {}
    names = [Path(d).name for d in change_dirs]
    titles = {}
    for change_dir in map(Path, change_dirs):
        name = change_dir.name
        proposal = _read(change_dir / "proposal.md")
        design = _read(change_dir / "design.md")
        tasks = _read(change_dir / "tasks.md")
        specs = sorted(change_dir.glob("specs/**/spec.md"))
        missing[name] = [label for label, ok in (("proposal.md", proposal), ("design.md", design),
                                                 ("specs/*/spec.md", specs), ("tasks.md", tasks)) if not ok]
        why = _section(proposal, "Why")
        if why:
            whys.append((name, why))
        whats += _bullets(_section(proposal, "What Changes"))
        for title, body in _decisions(design):
            titles.setdefault(title, []).append(name)
            decisions.append({"process": title, "contents": body, "_source": name})
        for spec in specs:
            scenarios += _scenarios(_read(spec))
        steps += _verification_steps(tasks)

    for d in decisions:  # prefix only when two changes share a decision title
        if len(set(titles[d["process"]])) > 1:
            d["process"] = f"{d['_source']}: {d['process']}"
        del d["_source"]

    if len(whys) == 1:
        why_body = whys[0][1]
    else:
        why_body = "\n".join(f"• {n}: {w}" for n, w in whys)
    how_body = "\n".join(f"{i}. {d['process']}" for i, d in enumerate(decisions, start=1))
    test_body = "Conditions:\n" + "\n".join(f"• {s}" for s in _dedupe(scenarios))
    test_body += "\n\nSteps:\n" + "\n".join(f"{i}. {s}" for i, s in enumerate(_dedupe(steps), start=1))
    suffix = names[0] if names else ""

    return {
        "meta": {"title": suffix, "outputSuffix": suffix, "sources": names, "missingArtifacts": missing},
        "overview": {"sections": [
            {"title": "1. Why the change", "body": why_body},
            {"title": "2. What is changed", "body": "\n".join(f"• {w}" for w in _dedupe(whats))},
            {"title": "3. How the process works", "body": how_body},
            {"title": "4. How to test", "body": test_body},
        ]},
        "uiux": {"note": "No screenshots yet. Insert UI screenshots below each caption.", "images": []},
        "diagrams": {"title": f"{suffix} - Process Flow", "instruction": DRAWIO_NOTE,
                     "mermaid": "flowchart TD\n    A[Start] --> B[...]", "renderedImage": ""},
        "detailedDesign": {"decisions": decisions, "apiStructure": {"present": False}},
        "databaseDesign": {"present": False, "tables": []},
    }


# --------------------------------------------------------------------------- cli

def fail(message):
    print(f"\n[ERROR] {message}")
    sys.exit(1)


def main(argv=None):
    parser = argparse.ArgumentParser(
        description="Build the change documentation workbook from reviewed JSON, "
                    "or draft that JSON from OpenSpec change folders.")
    parser.add_argument("--input", help="Reviewed JSON (work/temp_doc.json)")
    parser.add_argument("--output", help="Workbook path (work/Design Docs/<suffix> - Documentation.xlsx)")
    parser.add_argument("--draft", nargs="+", metavar="CHANGE_DIR", help="Change folder(s) to draft JSON from")
    parser.add_argument("--json-out", help="Where --draft writes the JSON")
    args = parser.parse_args(argv)

    if args.draft:
        if not args.json_out:
            fail("--draft needs --json-out")
        for d in args.draft:
            if not os.path.isdir(d):
                fail(f"Change folder not found: {d}")
        out = Path(args.json_out)
        if not out.parent.is_dir():
            fail(f"Output directory does not exist: {out.parent.resolve()}")
        draft = draft_from_changes(args.draft)
        out.write_text(json.dumps(draft, indent=2, ensure_ascii=False), encoding="utf-8")
        print(f"Draft written: {out}")
        for name, gaps in draft["meta"]["missingArtifacts"].items():
            if gaps:
                print(f"[WARN] {name}: missing {', '.join(gaps)}")
        return

    if not (args.input and args.output):
        parser.error("--input and --output are required (or use --draft ... --json-out ...)")
    if not os.path.exists(args.input):
        fail(f"Input file not found: {args.input}")
    try:
        with open(args.input, encoding="utf-8") as fh:
            data = json.load(fh)
    except Exception as exc:
        fail(f"Cannot read input JSON: {exc}")
    out_dir = os.path.dirname(os.path.abspath(args.output))
    if not os.path.isdir(out_dir):
        fail(f"Output directory does not exist: {out_dir}")

    wb = build_workbook(data)
    try:
        wb.save(args.output)
    except Exception as exc:
        fail(f"Cannot save workbook (is it open in Excel?): {exc}")
    print(f"Saved: {args.output}")
    print(f"Sheets: {wb.sheetnames}")


if __name__ == "__main__":
    main()
