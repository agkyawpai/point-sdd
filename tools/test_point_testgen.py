"""Run: .venv\\Scripts\\python -m unittest tools/test_point_testgen.py"""

import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

TOOLS = Path(__file__).resolve().parent
TESTGEN = TOOLS / "point-testgen.py"
MAKE_TEMPLATE = TOOLS / "point-make-template.py"

METHOD = "Scope: browser\nMutation: read-only\nPreconditions: x\nSteps:\n1. y\nCleanup: None"


def case(i, **extra):
    return {"selected": True, "Classification": "Normal", "Content": f"Case {i}",
            "Test method": METHOD, "Test result": f"Result {i}", **extra}


class TestGen(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.dir = Path(tempfile.mkdtemp())
        cls.template = cls.dir / "template.xlsx"
        subprocess.run([sys.executable, str(MAKE_TEMPLATE), "--output", str(cls.template)],
                       check=True, capture_output=True)

    def generate(self, tests, meta=None, check=True):
        src = self.dir / "in.json"
        out = self.dir / "out.xlsx"
        src.write_text(json.dumps({"meta": meta or {}, "tests": tests}, ensure_ascii=False),
                       encoding="utf-8")
        result = subprocess.run([sys.executable, str(TESTGEN), "--input", str(src),
                                 "--template", str(self.template), "--output", str(out)],
                                capture_output=True, text=True, check=check)
        return result, out

    def test_template_has_all_sheets(self):
        wb = load_workbook(self.template)
        self.assertEqual(wb.sheetnames, ["Title", "TestCase", "Evidence", "NG Report", "NG_evidence", "List"])
        headers = [wb["TestCase"].cell(row=6, column=c).value for c in (1, 3, 8, 13, 24, 35, 37, 39, 41, 44)]
        self.assertEqual(headers, ["No", "Classification", "Content", "Test method", "Test results",
                                   "Tester", "Test Date", "Result", "NG Report", "Confirmed date"])

    def test_selected_cases_are_written_and_numbered(self):
        _, out = self.generate([case(1), case(2, selected=False), case(3)])
        ws = load_workbook(out)["TestCase"]
        self.assertEqual([ws.cell(row=r, column=1).value for r in (7, 8, 9)], [1, 2, None])
        self.assertEqual(ws["H7"].value, "Case 1")
        self.assertEqual(ws["H8"].value, "Case 3")
        self.assertEqual(ws["X8"].value, "Result 3")

    def test_more_cases_than_template_rows_get_merged(self):
        _, out = self.generate([case(i) for i in range(1, 61)])
        ws = load_workbook(out)["TestCase"]
        self.assertEqual(ws["A66"].value, 60)
        merged = {str(m) for m in ws.merged_cells.ranges}
        self.assertIn("M66:W66", merged)
        self.assertIn("X66:AH66", merged)

    def test_missing_field_is_rejected(self):
        bad = case(1)
        bad["Test result"] = " "
        result, _ = self.generate([bad], check=False)
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("missing required fields", result.stdout)

    def test_metadata_written(self):
        meta = {"systemName": "Point Barbershop System", "subSystemName": "Login", "author": "Tester A"}
        _, out = self.generate([case(1)], meta)
        wb = load_workbook(out)
        self.assertEqual(wb["TestCase"]["Q1"].value, "Point Barbershop System")
        self.assertEqual(wb["TestCase"]["Q2"].value, "Login")
        self.assertEqual(wb["TestCase"]["AP1"].value, "Tester A")
        self.assertEqual(wb["Title"]["P18"].value, "Tester A")
        self.assertEqual(wb["NG Report"]["Q2"].value, "Login")

    def test_dropdowns_cover_case_rows(self):
        _, out = self.generate([case(i) for i in range(1, 6)])
        ws = load_workbook(out)["TestCase"]
        ranges = {dv.formula1: str(dv.sqref) for dv in ws.data_validations.dataValidation}
        self.assertEqual(ranges["List!$B$1:$B$20"], "AI7:AI11")
        self.assertEqual(ranges["List!$A$1:$A$10"], "AM7:AM11")


if __name__ == "__main__":
    unittest.main()
