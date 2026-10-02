"""Fill the test case template from a reviewed temp_tests.json.

    .venv\\Scripts\\python tools\\point-testgen.py --input <temp_tests.json> \\
        --template "templates\\Test Case - Function.xlsx" --output "work\\Test Cases\\Test Case - <suffix>.xlsx"

The output is a copy of the template with one row per selected case on the
TestCase sheet (from row 7), plus system / sub system / date / author written
to the Title, TestCase and NG Report sheets.
"""

import argparse
import json
import math
import shutil
import sys
from copy import copy
from datetime import date
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

REQUIRED = ("Classification", "Content", "Test method", "Test result")

START_ROW = 7          # first case row on TestCase
COL_NO = 1             # A
COL_CLASS = 3          # C
COL_CONTENT = 8        # H
COL_METHOD = 13        # M
COL_RESULT = 24        # X  ("Test results")
COL_TESTER = 35        # AI
COL_DATE = 37          # AK
COL_OK_NG = 39         # AM ("Result")
VALUE_COLUMNS = (COL_NO, COL_CLASS, COL_CONTENT, COL_METHOD, COL_RESULT,
                 COL_TESTER, COL_DATE, COL_OK_NG)
DEFAULT_WIDTH = 8.54


def fail(message):
    print(f"[ERROR] {message}")
    sys.exit(1)


def load_tests(path):
    try:
        data = json.loads(Path(path).read_text(encoding="utf-8"))
    except Exception as e:
        fail(f"Error reading JSON: {e}")
    if "tests" not in data:
        fail("Invalid JSON format: missing 'tests' key.")
    selected = [t for t in data["tests"] if t.get("selected", True)]
    for i, test in enumerate(selected, start=1):
        missing = [f for f in REQUIRED if not str(test.get(f, "")).strip()]
        if missing:
            fail(f"Test case {i} is missing required fields: {', '.join(missing)}")
    return data.get("meta", {}), selected


def merge_pattern(ws, row):
    """Column spans merged on a single template row, e.g. [(1, 2), (3, 7), ...]."""
    return sorted((m.min_col, m.max_col) for m in ws.merged_cells.ranges
                  if m.min_row == row and m.max_row == row)


def span_width(ws, c1, c2):
    total = 0.0
    for c in range(c1, c2 + 1):
        dim = ws.column_dimensions.get(get_column_letter(c))
        total += dim.width if dim is not None and dim.width else DEFAULT_WIDTH
    return total


def needed_height(ws, values, pattern, base_height):
    """Row height that fits the longest wrapped text in the row (merged cells don't auto-fit)."""
    spans = {c1: c2 for c1, c2 in pattern}
    lines = 1
    for col, text in values.items():
        if not isinstance(text, str) or col not in spans:
            continue
        chars = max(int(span_width(ws, col, spans[col]) * 1.15), 4)
        n = sum(max(1, math.ceil(len(part) / chars)) for part in text.split("\n"))
        lines = max(lines, n)
    return min(max(base_height, lines * 12 + 6), 409)


def write_cases(ws, tests):
    pattern = merge_pattern(ws, START_ROW)
    if not pattern:
        fail("Template TestCase sheet has no merged cells on row 7 — rebuild the template")
    base_height = ws.row_dimensions[START_ROW].height or 60
    existing = {(m.min_row, m.min_col) for m in ws.merged_cells.ranges}
    last_row = START_ROW + len(tests) - 1

    for i, test in enumerate(tests, start=1):
        row = START_ROW + i - 1
        values = {
            COL_NO: i,
            COL_CLASS: test["Classification"],
            COL_CONTENT: test["Content"],
            COL_METHOD: test["Test method"],
            COL_RESULT: test["Test result"],
        }
        # Rows beyond the template's pre-formatted rows get the same merges and styles.
        for c1, c2 in pattern:
            if (row, c1) not in existing:
                for c in range(c1, c2 + 1):
                    src, dst = ws.cell(row=START_ROW, column=c), ws.cell(row=row, column=c)
                    dst.font, dst.border = copy(src.font), copy(src.border)
                    dst.fill, dst.alignment = copy(src.fill), copy(src.alignment)
                    dst.number_format = src.number_format
                ws.merge_cells(start_row=row, start_column=c1, end_row=row, end_column=c2)
        for col, value in values.items():
            ws.cell(row=row, column=col, value=value)
        ws.row_dimensions[row].height = needed_height(ws, values, pattern, base_height)

    # Clear values in pre-formatted rows that no case uses, so they can't be read as cases.
    for row in range(last_row + 1, ws.max_row + 1):
        for col in VALUE_COLUMNS:
            ws.cell(row=row, column=col).value = None
    return last_row


def add_dropdowns(ws, last_row):
    ws.data_validations.dataValidation = []
    tester = DataValidation(type="list", formula1="List!$B$1:$B$20", allow_blank=True)
    tester.add(f"{get_column_letter(COL_TESTER)}{START_ROW}:{get_column_letter(COL_TESTER)}{last_row}")
    result = DataValidation(type="list", formula1="List!$A$1:$A$10", allow_blank=True)
    result.add(f"{get_column_letter(COL_OK_NG)}{START_ROW}:{get_column_letter(COL_OK_NG)}{last_row}")
    ws.add_data_validation(tester)
    ws.add_data_validation(result)


def write_meta(wb, meta):
    system = meta.get("systemName", "")
    sub = meta.get("subSystemName", "")
    author = meta.get("author", "")
    today = date.today()

    title = wb["Title"]
    if system:
        title["AU1"] = system
    for cell in ("L18", "L19"):
        title[cell] = today
    for cell in ("P18", "P19"):
        title[cell] = author

    for name in ("TestCase", "NG Report"):
        ws = wb[name]
        if system:
            ws["Q1"] = system
        ws["Q2"] = sub
        ws["AI1"] = today
        ws["AP1"] = author


def generate(input_json, template, output):
    meta, tests = load_tests(input_json)
    if not tests:
        print("No tests selected (all 'selected' are false or the list is empty). Nothing written.")
        return
    if not Path(template).exists():
        fail(f"Template not found: {template}")
    Path(output).parent.mkdir(parents=True, exist_ok=True)
    try:
        shutil.copyfile(template, output)
    except PermissionError:
        fail(f"Cannot write {output} — close it in Excel and run again")
    print(f"Created new file from template: {output}")

    wb = load_workbook(output)
    ws = wb["TestCase"]
    last_row = write_cases(ws, tests)
    add_dropdowns(ws, last_row)
    write_meta(wb, meta)
    try:
        wb.save(output)
    except PermissionError:
        fail(f"Cannot write {output} — close it in Excel and run again")
    print(f"Successfully wrote {len(tests)} test cases to {output}")


def main():
    ap = argparse.ArgumentParser(description="Generate an Excel test case workbook from JSON.")
    ap.add_argument("--input", required=True, help="temp_tests.json")
    ap.add_argument("--template", required=True, help="Template workbook")
    ap.add_argument("--output", required=True, help="Output workbook")
    args = ap.parse_args()
    generate(args.input, args.template, args.output)


if __name__ == "__main__":
    main()
