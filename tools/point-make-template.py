"""Create templates/Test Case - Function.xlsx — the blank workbook point-testgen.py fills.

    .venv\\Scripts\\python tools\\point-make-template.py [--output templates\\Test Case - Function.xlsx]

Edit this script (not the .xlsx) to change the layout, then re-run it.
Tester names for the dropdown are in the List sheet, column B — edit TESTERS below.
"""

import argparse
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import column_index_from_string as col_idx
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
DEFAULT_OUT = ROOT / "templates" / "Test Case - Function.xlsx"

SYSTEMS = ["Point Barbershop System", "Point Barbershop Website"]
TESTERS = ["Aung Kyaw Paing"]
RESULTS = ["OK", "NG"]
NG_TYPES = ["Specification leak", "Program miss", "Other"]

CASE_ROWS = range(7, 46)   # pre-formatted case rows on TestCase
NG_ROWS = range(7, 57)     # pre-formatted rows on NG Report
EVIDENCE_ROWS = 1500

FONT = "Arial"
F = Font(name=FONT, size=9)
F_TITLE = Font(name=FONT, size=14, bold=True)
F_HEAD = Font(name=FONT, size=9, bold=True)
GREEN = PatternFill("solid", fgColor="92D050")
THIN = Side(style="thin", color="000000")
MEDIUM = Side(style="medium", color="000000")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
RIGHT = Alignment(horizontal="right", vertical="center")

# TestCase columns: (header, first col, last col, alignment)
CASE_COLUMNS = [
    ("No", "A", "B", CENTER),
    ("Classification", "C", "G", CENTER),
    ("Content", "H", "L", LEFT),
    ("Test method", "M", "W", LEFT),
    ("Test results", "X", "AH", LEFT),
    ("Tester", "AI", "AJ", CENTER),
    ("Test Date", "AK", "AL", CENTER),
    ("Result", "AM", "AN", CENTER),
    ("NG Report", "AO", "AQ", CENTER),
    ("Confirmed date", "AR", "AS", CENTER),
]
NG_COLUMNS = [
    ("No", "A", "B", CENTER),
    ("Classification", "C", "G", CENTER),
    ("NG Content", "H", "W", LEFT),
    ("Modification Content", "X", "AM", LEFT),
    ("Modification Name", "AN", "AP", CENTER),
    ("Confirmation Name", "AQ", "AS", CENTER),
]
CASE_WIDTHS = {"A": 2.2, "B": 1.2, "C": 2.5, "G": 11.8, "H": 2.5, "L": 7.5, "M": 3.2, "O": 3.5,
               "Q": 3.8, "R": 3.5, "S": 3.2, "T": 3.8, "W": 18.2, "X": 4.2, "AH": 8.6, "AI": 3.0,
               "AJ": 8.5, "AK": 3.5, "AL": 6.6, "AM": 2.5, "AN": 6.5, "AO": 2.5, "AR": 3.5,
               "AS": 3.8, "AT": 8.5}


def box(ws, rng, value=None, font=F, align=CENTER, fill=None, border=BOX):
    """Merge a range, style every cell in it, and put value in the top-left."""
    ws.merge_cells(rng)
    for row in ws[rng]:
        for c in row:
            c.border = border
            c.font = font
            c.alignment = align
            if fill:
                c.fill = fill
    if value is not None:
        ws[rng.split(":")[0]] = value


def setup_sheet(ws, default_width=8.54):
    ws.sheet_format.defaultColWidth = default_width
    ws.sheet_format.defaultRowHeight = 16
    for row in ws.iter_rows(min_row=1, max_row=1):
        for c in row:
            c.font = F


def header_block(ws, title):
    """Rows 1–2: title, system / sub system, create date, author."""
    box(ws, "A1:J2", title, F_TITLE, CENTER, border=Border(left=MEDIUM, top=MEDIUM, bottom=MEDIUM, right=THIN))
    box(ws, "K1:P1", "System Name")
    box(ws, "Q1:AE1", "=Title!AU1")
    box(ws, "K2:P2", "Sub System Name")
    box(ws, "Q2:AE2")
    box(ws, "AF1:AH2", "Create\nDate")
    box(ws, "AI1:AL2", align=CENTER)
    ws["AI1"].number_format = "yyyy/mm/dd"
    box(ws, "AM1:AO2", "Author")
    box(ws, "AP1:AS2")
    for r in (1, 2):
        ws.row_dimensions[r].height = 16


def summary_cell(ws, label_col, label, value_rng, formula, unit="Case"):
    ws[f"{label_col}4"] = label
    ws[f"{label_col}4"].font = F
    ws[f"{label_col}4"].alignment = RIGHT
    colon = ws.cell(row=4, column=col_idx(label_col) + 1, value=":")
    colon.font, colon.alignment = F, CENTER
    ws.merge_cells(value_rng)
    first = ws[value_rng.split(":")[0]]
    first.value, first.font, first.alignment = formula, Font(name=FONT, size=9, bold=True), RIGHT
    end_col = col_idx("".join(ch for ch in value_rng.split(":")[1] if ch.isalpha()))
    u = ws.cell(row=4, column=end_col + 1, value=unit)
    u.font = F


def table(ws, columns, rows, header_row=6):
    for header, c1, c2, _ in columns:
        box(ws, f"{c1}{header_row}:{c2}{header_row}", header, F_HEAD, CENTER, GREEN)
    for r in rows:
        for _, c1, c2, align in columns:
            box(ws, f"{c1}{r}:{c2}{r}", align=align)


# ── Sheets ─────────────────────────────────────────────────────────────────

def title_sheet(wb):
    ws = wb.active
    ws.title = "Title"
    setup_sheet(ws)
    ws.column_dimensions["A"].width = 2.5
    ws.column_dimensions["AU"].width = 23.5
    for i, name in enumerate(SYSTEMS, start=1):
        ws[f"AT{i}"] = name
    ws["AU1"] = SYSTEMS[0]           # current system name; TestCase!Q1 reads it
    for c in ("AT1", "AT2", "AU1"):
        ws[c].font = Font(name=FONT, size=9, color="808080")
    box(ws, "I13:AK14", "Unit Test Specification and Report", Font(name=FONT, size=20, bold=True),
        CENTER, border=Border(bottom=MEDIUM))
    labels = [("H", "K", ""), ("L", "O", "Created date"), ("P", "T", "Author"),
              ("U", "X", "Review date"), ("Y", "AC", "Reviewer"), ("AD", "AG", "Approval date"),
              ("AH", "AL", "Authorizer")]
    for c1, c2, text in labels:
        box(ws, f"{c1}17:{c2}17", text or None, F_HEAD, CENTER, GREEN)
    for r, label in ((18, "Specification"), (19, "Report")):
        for c1, c2, _ in labels:
            box(ws, f"{c1}{r}:{c2}{r}", label if c1 == "H" else None)
        for c in ("L", "U", "AD"):
            ws[f"{c}{r}"].number_format = "yyyy/mm/dd"
    for r in (13, 14, 17, 18, 19):
        ws.row_dimensions[r].height = 16


def testcase_sheet(wb):
    ws = wb.create_sheet("TestCase")
    setup_sheet(ws)
    for col, w in CASE_WIDTHS.items():
        ws.column_dimensions[col].width = w
    header_block(ws, "Unit Test Specifications")
    summary_cell(ws, "G", "Total number of test cases", "I4:K4", "=COUNTA(X7:AH1048576)")
    summary_cell(ws, "R", "OK", "T4:V4", '=COUNTIF(AM7:AN1048576,"OK")')
    summary_cell(ws, "AB", "NG", "AD4:AF4", '=COUNTIF(AM7:AN1048576,"NG")')
    summary_cell(ws, "AL", "Progress", "AN4:AP4", "=IFERROR(T4/I4*100,0)", unit="%")
    ws["AN4"].number_format = "0.0"
    table(ws, CASE_COLUMNS, CASE_ROWS)
    for r in CASE_ROWS:
        ws.row_dimensions[r].height = 60.9
        ws[f"AK{r}"].number_format = "yyyy/mm/dd"
        ws[f"AR{r}"].number_format = "yyyy/mm/dd"
    last = CASE_ROWS[-1]
    tester = DataValidation(type="list", formula1="List!$B$1:$B$20", allow_blank=True)
    tester.add(f"AI7:AI{last}")
    result = DataValidation(type="list", formula1="List!$A$1:$A$10", allow_blank=True)
    result.add(f"AM7:AM{last}")
    ws.add_data_validation(tester)
    ws.add_data_validation(result)
    ws.freeze_panes = "A7"
    ws.sheet_view.zoomScale = 74


def evidence_sheet(wb):
    ws = wb.create_sheet("Evidence")
    ws.column_dimensions["A"].width = 8.9
    ws["A1"], ws["B1"] = "No", "Please paste the screenshot below this. If necessary, attach data etc."
    ws["A1"].font = ws["B1"].font = F_HEAD
    ws["A2"] = 1
    for r in range(3, EVIDENCE_ROWS + 1):
        ws[f"A{r}"] = '=IF(MOD(ROW(),15)=0,((ROW())/15)+1,"")'
    for r in range(2, EVIDENCE_ROWS + 1):
        ws[f"A{r}"].font = F_HEAD
        ws[f"A{r}"].alignment = CENTER
    ws.freeze_panes = "A2"


def ng_report_sheet(wb):
    ws = wb.create_sheet("NG Report")
    setup_sheet(ws)
    ws.column_dimensions["A"].width = 2.5
    header_block(ws, "Unit Test Specifications")
    summary_cell(ws, "G", "Total number of modifications", "I4:K4", "=COUNTA(H7:W1048576)")
    summary_cell(ws, "R", "Specification leak", "T4:V4", '=COUNTIF($C7:$G1048576,"Specification leak")')
    summary_cell(ws, "AB", "Program miss", "AD4:AF4", '=COUNTIF($C7:$G1048576,"Program miss")')
    summary_cell(ws, "AL", "Other", "AN4:AP4", '=COUNTIF($C7:$G1048576,"Other")')
    table(ws, NG_COLUMNS, NG_ROWS)
    for r in NG_ROWS:
        ws[f"A{r}"] = "=ROW()-6"
        ws.row_dimensions[r].height = 18
    dv = DataValidation(type="list", formula1='"' + ",".join(NG_TYPES) + '"', allow_blank=True)
    dv.add(f"C{NG_ROWS[0]}:G{NG_ROWS[-1]}")
    ws.add_data_validation(dv)
    ws.freeze_panes = "A7"
    ws.sheet_view.zoomScale = 130


def ng_evidence_sheet(wb):
    ws = wb.create_sheet("NG_evidence")
    ws.column_dimensions["A"].width = 9.2
    ws["A1"], ws["C1"] = "No", "Please paste the screenshot below this. If necessary, attach data etc."
    ws["A1"].font = ws["C1"].font = F_HEAD
    ws["A3"] = 1
    ws["A3"].font = F_HEAD
    ws.freeze_panes = "A2"


def list_sheet(wb):
    ws = wb.create_sheet("List")
    for i, v in enumerate(RESULTS, start=1):
        ws.cell(row=i, column=1, value=v).font = F
    for i, v in enumerate(TESTERS, start=1):
        ws.cell(row=i, column=2, value=v).font = F


def main():
    ap = argparse.ArgumentParser(description="Create the blank test case template.")
    ap.add_argument("--output", type=Path, default=DEFAULT_OUT)
    args = ap.parse_args()
    wb = Workbook()
    title_sheet(wb)
    testcase_sheet(wb)
    evidence_sheet(wb)
    ng_report_sheet(wb)
    ng_evidence_sheet(wb)
    list_sheet(wb)
    wb.active = 1
    args.output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.output)
    print(f"Wrote {args.output}")


if __name__ == "__main__":
    main()
