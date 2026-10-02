import hashlib
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl

SCRIPT_DIR = Path(__file__).resolve().parents[1] / "scripts"
sys.path.insert(0, str(SCRIPT_DIR))

from workbook_tool import (  # noqa: E402
    WorkbookError,
    apply_results,
    inspect_workbook,
    parse_method,
    validate_results,
)

# Header layout of the generated TestCase sheet: (label, first column, last column).
LAYOUT = [
    ("No", "A", "B"),
    ("Classification", "C", "G"),
    ("Content", "H", "L"),
    ("Test method", "M", "W"),
    ("Test results", "X", "AH"),
    ("Tester", "AI", "AJ"),
    ("Test Date", "AK", "AL"),
    ("Result", "AM", "AN"),
    ("NG Report", "AO", "AQ"),
    ("Confirmed date", "AR", "AS"),
]
HEADER_ROW = 6

CASES = [
    {
        "no": 1,
        "classification": "Normal",
        "content": "Customer list shows active customers",
        "method": "Scope: browser\nMutation: read-only\nPreconditions: signed in as admin\nSteps:\n1. Open Customers\nCleanup: none",
        "expected": "Active customers are listed",
    },
    {
        "no": 2,
        "classification": "Abnormal",
        "content": "Booking with no barber is rejected",
        "method": "Scope: browser-context\nMutation: data-mutating\nPreconditions: none\nSteps:\n1. Submit empty form\nCleanup: delete draft",
        "expected": "Validation error is shown",
        "result": "OK",
        "tester": "Existing tester",
    },
    {
        "no": 3,
        "classification": "Boundary",
        "content": "Service name accepts 100 characters",
        "method": "Scope: ambiguous\nMutation: RBAC-mutating\nSteps:\n1. Enter 100 characters",
        "expected": "Name is saved",
    },
]


def build_workbook(path: Path, *, layout=LAYOUT, header_row=HEADER_ROW, sheet_name="TestCase", cases=CASES) -> Path:
    workbook = openpyxl.Workbook()
    title = workbook.active
    title.title = "Title"
    title["A1"] = "Point Barbershop System - Test Specification"
    sheet = workbook.create_sheet(sheet_name)
    sheet["A2"] = "Function: Customers"

    def place(row, label, value):
        first = next(f for name, f, _ in layout if name == label)
        last = next(l for name, _, l in layout if name == label)
        if first != last:
            sheet.merge_cells(f"{first}{row}:{last}{row}")
        if value is not None:
            sheet[f"{first}{row}"] = value

    for label, _, _ in layout:
        place(header_row, label, label)
    for offset, case in enumerate(cases, start=1):
        row = header_row + offset
        place(row, "No", case["no"])
        place(row, "Classification", case.get("classification"))
        place(row, "Content", case.get("content"))
        place(row, "Test method", case.get("method"))
        place(row, "Test results", case.get("expected"))
        place(row, "Tester", case.get("tester"))
        place(row, "Test Date", case.get("date"))
        place(row, "Result", case.get("result"))
        place(row, "NG Report", None)
        place(row, "Confirmed date", None)
    for name in ("Evidence", "NG Report", "NG_evidence"):
        workbook.create_sheet(name)
    lists = workbook.create_sheet("List")
    lists["A1"], lists["A2"] = "OK", "NG"
    lists["B1"], lists["B2"] = "QA One", "QA Two"
    workbook.save(path)
    return path


def results_payload(cases, **extra):
    payload = {
        "schemaVersion": 1,
        "runId": "20261002T000000Z-customers-abc123",
        "workbook": "Test Case - Customers.xlsx",
        "completedAt": "2026-10-02T00:00:00Z",
        "cases": cases,
    }
    payload.update(extra)
    return payload


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


class InspectTests(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.root = Path(self._tmp.name)
        self.workbook = build_workbook(self.root / "Test Case - Customers.xlsx")

    def tearDown(self):
        self._tmp.cleanup()

    def test_discovers_sheet_header_and_merged_columns(self):
        manifest = inspect_workbook(self.workbook)
        self.assertEqual(manifest["sheet"], "TestCase")
        self.assertEqual(manifest["headerRow"], 6)
        self.assertEqual(manifest["firstDataRow"], 7)
        self.assertEqual(
            manifest["columns"],
            {
                "caseNo": "A",
                "classification": "C",
                "content": "H",
                "method": "M",
                "expected": "X",
                "tester": "AI",
                "date": "AK",
                "result": "AM",
                "ngReport": "AO",
                "confirmedDate": "AR",
            },
        )
        self.assertEqual(manifest["headerSpans"]["method"], "M6:W6")
        self.assertIn("NG_evidence", manifest["sheets"])

    def test_reads_cases_with_scope_and_mutation(self):
        cases = inspect_workbook(self.workbook)["cases"]
        self.assertEqual([case["caseNo"] for case in cases], [1, 2, 3])
        self.assertEqual([case["row"] for case in cases], [7, 8, 9])
        first, second, third = cases
        self.assertEqual(first["classification"], "Normal")
        self.assertEqual((first["scope"], first["mutation"]), ("browser", "read-only"))
        self.assertEqual((second["scope"], second["mutation"]), ("browser-context", "data-mutating"))
        self.assertEqual(second["result"], "OK")
        self.assertEqual((third["scope"], third["mutation"]), ("ambiguous", "RBAC-mutating"))

    def test_reads_list_sheet(self):
        lists = inspect_workbook(self.workbook)["lists"]
        self.assertEqual(lists["resultValues"], ["OK", "NG"])
        self.assertEqual(lists["testers"], ["QA One", "QA Two"])

    def test_does_not_assume_fixed_positions(self):
        shifted = [
            ("Result", "B", "C"),
            ("Test method", "D", "D"),
            ("No", "E", "E"),
            ("Test results", "F", "H"),
            ("Content", "I", "I"),
            ("Tester", "J", "J"),
            ("Test Date", "K", "K"),
            ("Classification", "L", "L"),
            ("NG Report", "M", "M"),
            ("Confirmed date", "N", "N"),
        ]
        path = build_workbook(self.root / "shifted.xlsx", layout=shifted, header_row=3, sheet_name="Cases")
        manifest = inspect_workbook(path)
        self.assertEqual(manifest["sheet"], "Cases")
        self.assertEqual(manifest["headerRow"], 3)
        self.assertEqual(manifest["columns"]["caseNo"], "E")
        self.assertEqual(manifest["columns"]["result"], "B")
        self.assertEqual([case["row"] for case in manifest["cases"]], [4, 5, 6])

    def test_skips_rows_without_case_content(self):
        cases = CASES + [{"no": 4}, {"no": "not a number", "content": "Note row"}]
        path = build_workbook(self.root / "sparse.xlsx", cases=cases)
        self.assertEqual([case["caseNo"] for case in inspect_workbook(path)["cases"]], [1, 2, 3])

    def test_rejects_workbook_without_headers(self):
        path = self.root / "empty.xlsx"
        openpyxl.Workbook().save(path)
        with self.assertRaises(WorkbookError):
            inspect_workbook(path)

    def test_cli_inspect_prints_json(self):
        completed = subprocess.run(
            [sys.executable, str(SCRIPT_DIR / "workbook_tool.py"), "inspect", str(self.workbook)],
            capture_output=True,
            text=True,
            encoding="utf-8",
            check=True,
        )
        self.assertEqual(json.loads(completed.stdout)["columns"]["result"], "AM")

    def test_parse_method(self):
        self.assertEqual(parse_method("Scope: browser\nMutation: read-only\nSteps:\n1. x"), {"scope": "browser", "mutation": "read-only"})
        self.assertEqual(parse_method(None), {})


class ApplyResultsTests(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.root = Path(self._tmp.name)
        self.workbook = build_workbook(self.root / "Test Case - Customers.xlsx")
        self.run_dir = self.root / "Customers Evidence" / "runs" / "run-1"

    def tearDown(self):
        self._tmp.cleanup()

    def write_results(self, payload) -> Path:
        path = self.root / "results.json"
        path.write_text(json.dumps(payload), encoding="utf-8")
        return path

    def sheet(self):
        return openpyxl.load_workbook(self.workbook, data_only=False)["TestCase"]

    def test_maps_pass_fail_and_leaves_other_statuses_alone(self):
        results = self.write_results(
            results_payload(
                [
                    {"caseNo": 1, "status": "PASS", "tester": "QA One", "date": "2026/10/02", "evidence": ["screenshots/TC01-list.png"]},
                    {"caseNo": 2, "status": "BLOCKED", "tester": "QA One", "date": "2026/10/02"},
                    {"caseNo": 3, "status": "FAIL", "tester": "QA Two", "date": "2026/10/02"},
                ]
            )
        )
        report = apply_results(self.workbook, results, self.run_dir)
        sheet = self.sheet()
        self.assertEqual((sheet["AM7"].value, sheet["AI7"].value, sheet["AK7"].value), ("OK", "QA One", "2026/10/02"))
        self.assertEqual((sheet["AM8"].value, sheet["AI8"].value, sheet["AK8"].value), ("OK", "Existing tester", None))
        self.assertEqual(sheet["AM9"].value, "NG")
        self.assertEqual(report["totals"], {"OK": 1, "NG": 1})
        self.assertEqual(report["skipped"], [{"caseNo": 2, "status": "BLOCKED"}])
        self.assertTrue((self.run_dir / "Test Case - Customers-before.xlsx").is_file())
        self.assertIn("AM7:AN7", {str(r) for r in sheet.merged_cells.ranges})

    def test_out_of_scope_and_not_run_do_not_touch_rows(self):
        before = self.sheet()
        snapshot = [(cell.coordinate, cell.value) for row in before.iter_rows(min_row=7, max_row=9) for cell in row]
        results = self.write_results(
            results_payload([{"caseNo": 1, "status": "OUT_OF_SCOPE", "tester": "QA"}, {"caseNo": 3, "status": "NOT_RUN"}])
        )
        apply_results(self.workbook, results, self.run_dir)
        after = self.sheet()
        self.assertEqual(snapshot, [(cell.coordinate, cell.value) for row in after.iter_rows(min_row=7, max_row=9) for cell in row])

    def test_writes_to_merge_anchor_for_inner_cells(self):
        from openpyxl.utils import column_index_from_string
        from workbook_tool import _merge_anchors, _write

        sheet = openpyxl.load_workbook(self.workbook)["TestCase"]
        anchors = _merge_anchors(sheet)
        _write(sheet, anchors, 8, column_index_from_string("AN"), "NG")  # AN8 is inside AM8:AN8
        self.assertEqual(sheet["AM8"].value, "NG")

    def test_invalid_results_leave_workbook_unchanged(self):
        bad_payloads = [
            results_payload([{"caseNo": 1, "status": "PASS"}, {"caseNo": 1, "status": "FAIL"}]),
            results_payload([{"caseNo": 1, "status": "PASS", "evidence": ["C:\\evidence\\TC01.png"]}]),
            results_payload([{"caseNo": 1, "status": "PASS", "evidence": ["../other/TC01.png"]}]),
            results_payload([{"caseNo": 1, "status": "OK"}]),
            results_payload([{"caseNo": 1, "status": "PASS", "password": "x"}]),
            results_payload([{"caseNo": True, "status": "PASS"}]),
            results_payload([{"caseNo": 99, "status": "PASS"}]),
            results_payload([{"caseNo": 1, "row": 30, "status": "PASS"}]),
            results_payload([{"caseNo": 1, "status": "PASS"}], workbook="Test Case - Other.xlsx"),
            results_payload([{"caseNo": 1, "status": "PASS"}], schemaVersion=2),
            results_payload([{"caseNo": 1, "status": "PASS"}], completedAt=None),
        ]
        original = digest(self.workbook)
        for payload in bad_payloads:
            with self.subTest(payload=payload["cases"]):
                results = self.write_results(payload)
                with self.assertRaises(WorkbookError):
                    apply_results(self.workbook, results, self.run_dir)
                self.assertEqual(digest(self.workbook), original)
        self.assertFalse(self.run_dir.exists(), "no backup should be written when validation fails")

    def test_incomplete_run_can_be_allowed_explicitly(self):
        results = self.write_results(results_payload([{"caseNo": 1, "status": "PASS"}], completedAt=None))
        apply_results(self.workbook, results, self.run_dir, require_complete=False)
        self.assertEqual(self.sheet()["AM7"].value, "OK")

    def test_second_apply_keeps_first_backup(self):
        results = self.write_results(results_payload([{"caseNo": 1, "status": "PASS"}]))
        apply_results(self.workbook, results, self.run_dir)
        apply_results(self.workbook, results, self.run_dir)
        self.assertEqual(len(list(self.run_dir.glob("*-before*.xlsx"))), 2)

    def test_warns_when_tester_not_in_list(self):
        results = self.write_results(results_payload([{"caseNo": 1, "status": "PASS", "tester": "Somebody"}]))
        report = apply_results(self.workbook, results, self.run_dir)
        self.assertEqual(len(report["warnings"]), 1)


class ValidateResultsTests(unittest.TestCase):
    def test_accepts_all_statuses(self):
        cases = [{"caseNo": index, "status": status} for index, status in enumerate(["PASS", "FAIL", "BLOCKED", "OUT_OF_SCOPE", "NOT_RUN"], 1)]
        self.assertEqual(len(validate_results(results_payload(cases))["cases"]), 5)

    def test_rejects_non_list_cases(self):
        payload = results_payload({"caseNo": 1})
        with self.assertRaises(WorkbookError):
            validate_results(payload)


if __name__ == "__main__":
    unittest.main()
