"""Inspect Point QA test workbooks and apply validated browser-run results.

Commands:
  inspect <workbook>                          print a JSON manifest of the test-case sheet
  apply-results <workbook> <results.json> <backup-dir>
                                              write PASS/FAIL outcomes as OK/NG

The sheet, header row, case rows and columns are discovered from header text;
no fixed sheet layout or column letters are assumed.
"""

from __future__ import annotations

import argparse
import json
import os
import re
import shutil
import sys
import tempfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.utils import get_column_letter

SCHEMA_VERSION = 1
STATUSES = ("PASS", "FAIL", "BLOCKED", "OUT_OF_SCOPE", "NOT_RUN")
EXCEL_RESULT = {"PASS": "OK", "FAIL": "NG"}
HEADER_SCAN_ROWS = 60

# Normalized header text -> manifest key. The first matching column wins.
HEADER_ALIASES: dict[str, set[str]] = {
    "caseNo": {"no", "no.", "case", "case no", "case no.", "case number"},
    "classification": {"classification", "category", "type"},
    "content": {"content", "test case", "test content"},
    "method": {"test method", "method"},
    "expected": {"test results", "expected result", "expected results", "expected"},
    "tester": {"tester", "tested by"},
    "date": {"test date", "date", "tested date"},
    "result": {"result", "status"},
    "ngReport": {"ng report", "ng no", "ng no."},
    "confirmedDate": {"confirmed date", "confirmation date"},
}
REQUIRED_COLUMNS = {"caseNo", "content", "method", "expected", "result"}
FORBIDDEN_RESULT_KEYS = {"password", "passwd", "pwd", "token", "cookie", "cookies", "authorization", "credentials", "secret"}
METHOD_FIELDS = {"scope": "scope", "mutation": "mutation"}


class WorkbookError(ValueError):
    """Raised when a workbook or results file fails validation."""


def _normal(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value if value is not None else "")).strip().lower()


def _is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _merge_anchors(sheet) -> dict[tuple[int, int], tuple[int, int]]:
    """Map every cell inside a merged range to the range's top-left cell."""
    anchors: dict[tuple[int, int], tuple[int, int]] = {}
    for merged in sheet.merged_cells.ranges:
        for row in range(merged.min_row, merged.max_row + 1):
            for column in range(merged.min_col, merged.max_col + 1):
                anchors[(row, column)] = (merged.min_row, merged.min_col)
    return anchors


def _merge_span(sheet, row: int, column: int) -> str:
    for merged in sheet.merged_cells.ranges:
        if merged.min_row <= row <= merged.max_row and merged.min_col <= column <= merged.max_col:
            return merged.coord
    return f"{get_column_letter(column)}{row}"


def _value(sheet, anchors, row: int, column: int) -> Any:
    anchor_row, anchor_column = anchors.get((row, column), (row, column))
    return sheet.cell(anchor_row, anchor_column).value


def _find_header(sheet, anchors) -> tuple[int, dict[str, int]] | None:
    for row in range(1, min(sheet.max_row, HEADER_SCAN_ROWS) + 1):
        found: dict[str, int] = {}
        for column in range(1, sheet.max_column + 1):
            # Only the top-left cell of a merged header carries its label.
            if anchors.get((row, column), (row, column)) != (row, column):
                continue
            label = _normal(sheet.cell(row, column).value)
            if not label:
                continue
            for key, aliases in HEADER_ALIASES.items():
                if key not in found and label in aliases:
                    found[key] = column
                    break
        if REQUIRED_COLUMNS.issubset(found):
            return row, found
    return None


def _case_number(value: Any) -> int | None:
    if isinstance(value, bool) or _is_blank(value):
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        return int(value) if value.is_integer() else None
    text = str(value).strip()
    return int(text) if re.fullmatch(r"\d+", text) else None


def parse_method(text: Any) -> dict[str, str]:
    """Read the `Scope:` and `Mutation:` declarations at the top of a test method."""
    declared: dict[str, str] = {}
    for line in str(text or "").splitlines():
        match = re.match(r"\s*([A-Za-z]+)\s*:\s*(.*?)\s*$", line)
        if match and match.group(1).lower() in METHOD_FIELDS and match.group(2):
            declared.setdefault(METHOD_FIELDS[match.group(1).lower()], match.group(2))
    return declared


def _list_values(workbook) -> dict[str, list[str]]:
    sheet = next((workbook[name] for name in workbook.sheetnames if name.strip().lower() == "list"), None)
    if sheet is None:
        return {}
    values: dict[str, list[str]] = {"resultValues": [], "testers": []}
    for row in range(1, sheet.max_row + 1):
        for key, column in (("resultValues", 1), ("testers", 2)):
            cell = sheet.cell(row, column).value
            if not _is_blank(cell):
                values[key].append(str(cell).strip())
    return values


def _load(path: Path):
    if not path.is_file():
        raise FileNotFoundError(path)
    # Full (non read-only) mode is needed to see merged ranges.
    return openpyxl.load_workbook(path, data_only=False)


def _locate(workbook) -> tuple[Any, dict, int, dict[str, int]]:
    names = sorted(workbook.sheetnames, key=lambda name: _normal(name).replace(" ", "") != "testcase")
    for name in names:
        sheet = workbook[name]
        anchors = _merge_anchors(sheet)
        header = _find_header(sheet, anchors)
        if header:
            return sheet, anchors, header[0], header[1]
    raise WorkbookError("No sheet contains the required headers (No, Content, Test method, Test results, Result)")


def _manifest(path: Path, workbook) -> dict[str, Any]:
    sheet, anchors, header_row, columns = _locate(workbook)
    cases = []
    seen: dict[int, int] = {}
    for row in range(header_row + 1, sheet.max_row + 1):
        if anchors.get((row, columns["caseNo"]), (row, columns["caseNo"]))[0] != row:
            continue  # continuation of a vertically merged row
        case_no = _case_number(_value(sheet, anchors, row, columns["caseNo"]))
        if case_no is None:
            continue
        fields = {key: _value(sheet, anchors, row, column) for key, column in columns.items() if key != "caseNo"}
        if all(_is_blank(fields.get(key)) for key in ("content", "method", "expected")):
            continue
        if case_no in seen:
            raise WorkbookError(f"Case number {case_no} appears on rows {seen[case_no]} and {row}")
        seen[case_no] = row
        case = {"caseNo": case_no, "row": row, **fields}
        case.update(parse_method(fields.get("method")))
        cases.append(case)
    return {
        "schemaVersion": SCHEMA_VERSION,
        "workbook": str(path),
        "sheets": list(workbook.sheetnames),
        "sheet": sheet.title,
        "headerRow": header_row,
        "firstDataRow": cases[0]["row"] if cases else header_row + 1,
        "columns": {key: get_column_letter(column) for key, column in columns.items()},
        "headerSpans": {key: _merge_span(sheet, header_row, column) for key, column in columns.items()},
        "lists": _list_values(workbook),
        "cases": cases,
    }


def inspect_workbook(workbook_path: str | Path) -> dict[str, Any]:
    path = Path(workbook_path).resolve()
    workbook = _load(path)
    try:
        return _manifest(path, workbook)
    finally:
        workbook.close()


def _portable(value: Any) -> bool:
    if not isinstance(value, str) or not value.strip():
        return False
    text = value.strip().replace("\\", "/")
    return not (
        text.startswith(("/", "~"))
        or re.match(r"^[A-Za-z]:", text)
        or re.match(r"^[a-z]+://", text, re.IGNORECASE)
        or ".." in text.split("/")
    )


def validate_results(results: Any, *, require_complete: bool = True) -> dict[str, Any]:
    if not isinstance(results, dict):
        raise WorkbookError("results.json must contain an object")
    if results.get("schemaVersion") != SCHEMA_VERSION:
        raise WorkbookError(f"Unsupported results schema: {results.get('schemaVersion')!r}")
    if require_complete and not results.get("completedAt"):
        raise WorkbookError("results.json has no completedAt; finalize the run before applying it")
    cases = results.get("cases")
    if not isinstance(cases, list):
        raise WorkbookError("results.cases must be a list")
    seen: set[int] = set()
    for entry in cases:
        if not isinstance(entry, dict):
            raise WorkbookError("Each case result must be an object")
        case_no = entry.get("caseNo")
        if isinstance(case_no, bool) or not isinstance(case_no, int):
            raise WorkbookError(f"caseNo must be an integer: {case_no!r}")
        if case_no in seen:
            raise WorkbookError(f"Duplicate case number: {case_no}")
        seen.add(case_no)
        if entry.get("status") not in STATUSES:
            raise WorkbookError(f"Invalid status for case {case_no}: {entry.get('status')!r}")
        evidence = entry.get("evidence", [])
        if not isinstance(evidence, list):
            raise WorkbookError(f"evidence for case {case_no} must be a list")
        for item in evidence:
            if not _portable(item):
                raise WorkbookError(f"Evidence path for case {case_no} must be relative to the run directory: {item!r}")
        if FORBIDDEN_RESULT_KEYS.intersection(str(key).lower() for key in entry):
            raise WorkbookError(f"Credential-like fields are not allowed in case {case_no}")
        for key in ("tester", "date"):
            if key in entry and not isinstance(entry[key], str):
                raise WorkbookError(f"{key} for case {case_no} must be a string")
    return results


def _write(sheet, anchors, row: int, column: int, value: Any) -> None:
    anchor_row, anchor_column = anchors.get((row, column), (row, column))
    sheet.cell(anchor_row, anchor_column).value = value


def _backup(workbook_path: Path, backup_dir: Path) -> Path:
    backup_dir.mkdir(parents=True, exist_ok=True)
    target = backup_dir / f"{workbook_path.stem}-before.xlsx"
    if target.exists():
        stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
        target = backup_dir / f"{workbook_path.stem}-before-{stamp}.xlsx"
    shutil.copy2(workbook_path, target)
    return target


def apply_results(
    workbook_path: str | Path,
    results_path: str | Path,
    backup_dir: str | Path,
    *,
    require_complete: bool = True,
) -> dict[str, Any]:
    """Validate everything first, then write OK/NG (plus tester and date) for PASS/FAIL rows."""
    workbook_path = Path(workbook_path).resolve()
    results = validate_results(json.loads(Path(results_path).read_text(encoding="utf-8")), require_complete=require_complete)
    named = results.get("workbook")
    if named and Path(str(named)).name.lower() != workbook_path.name.lower():
        raise WorkbookError(f"results.json belongs to {named!r}, not {workbook_path.name!r}")

    workbook = _load(workbook_path)
    try:
        manifest = _manifest(workbook_path, workbook)
        by_case = {case["caseNo"]: case for case in manifest["cases"]}
        unknown = sorted(entry["caseNo"] for entry in results["cases"] if entry["caseNo"] not in by_case)
        if unknown:
            raise WorkbookError(f"Results contain case numbers not in the workbook: {unknown}")
        mismatched = [
            entry["caseNo"]
            for entry in results["cases"]
            if isinstance(entry.get("row"), int) and entry["row"] != by_case[entry["caseNo"]]["row"]
        ]
        if mismatched:
            raise WorkbookError(f"Result rows disagree with the workbook for cases {mismatched}; re-inspect the workbook")

        sheet = workbook[manifest["sheet"]]
        anchors = _merge_anchors(sheet)
        _, columns = _find_header(sheet, anchors)  # type: ignore[misc]
        testers = set(manifest["lists"].get("testers", []))
        applied, skipped, warnings = [], [], []
        for entry in results["cases"]:
            case_no, status = entry["caseNo"], entry["status"]
            if status not in EXCEL_RESULT:
                skipped.append({"caseNo": case_no, "status": status})
                continue
            row = by_case[case_no]["row"]
            _write(sheet, anchors, row, columns["result"], EXCEL_RESULT[status])
            if entry.get("tester") and "tester" in columns:
                _write(sheet, anchors, row, columns["tester"], entry["tester"])
                if testers and entry["tester"] not in testers:
                    warnings.append(f"Tester {entry['tester']!r} (case {case_no}) is not in the List sheet")
            if entry.get("date") and "date" in columns:
                _write(sheet, anchors, row, columns["date"], entry["date"])
            applied.append({"caseNo": case_no, "row": row, "result": EXCEL_RESULT[status]})
        backup = _backup(workbook_path, Path(backup_dir))

        handle, temporary = tempfile.mkstemp(prefix=f".{workbook_path.stem}-", suffix=workbook_path.suffix, dir=workbook_path.parent)
        os.close(handle)
        try:
            workbook.save(temporary)
            check = openpyxl.load_workbook(temporary, data_only=False)
            try:
                check_sheet = check[manifest["sheet"]]
                for item in applied:
                    if check_sheet.cell(item["row"], columns["result"]).value != item["result"]:
                        raise WorkbookError(f"Verification failed for case {item['caseNo']}")
            finally:
                check.close()
            try:
                os.replace(temporary, workbook_path)
            except PermissionError as error:
                raise WorkbookError(f"Cannot replace {workbook_path.name}; close it in Excel and retry") from error
        finally:
            if os.path.exists(temporary):
                os.unlink(temporary)
    finally:
        workbook.close()

    totals = {"OK": sum(1 for item in applied if item["result"] == "OK"), "NG": sum(1 for item in applied if item["result"] == "NG")}
    return {"workbook": str(workbook_path), "backup": str(backup), "applied": applied, "skipped": skipped, "totals": totals, "warnings": warnings}


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    commands = parser.add_subparsers(dest="command", required=True)
    inspect_parser = commands.add_parser("inspect", help="print the workbook manifest as JSON")
    inspect_parser.add_argument("workbook")
    apply_parser = commands.add_parser("apply-results", help="write OK/NG for PASS/FAIL results")
    apply_parser.add_argument("workbook")
    apply_parser.add_argument("results")
    apply_parser.add_argument("backup_dir", help="usually the run directory; receives a copy of the workbook before writing")
    args = parser.parse_args(argv)
    try:
        if args.command == "inspect":
            output = inspect_workbook(args.workbook)
        else:
            output = apply_results(args.workbook, args.results, args.backup_dir)
    except (WorkbookError, FileNotFoundError, json.JSONDecodeError) as error:
        print(f"error: {error}", file=sys.stderr)
        return 2
    print(json.dumps(output, ensure_ascii=False, indent=2, default=str))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
